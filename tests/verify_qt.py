# -*- coding: utf-8 -*-
"""功能验证: 默认27张 -> 全部加入 -> KPI/明细/合计/按商品汇总/CSV/预览"""
import io
import os
import sys
import time

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)                    # 仓库根目录
TMP = os.environ.get('TEMP') or HERE            # 临时目录(测试产物放这)
TEST_DIR = os.environ.get(                      # 测试发票目录(按需覆盖)
    'FAPIAO_TEST_DIR', r'C:\Users\yqh\Desktop\nj542发票')

sys.path.insert(0, ROOT)
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

os.environ["FAPIAO_HIST"] = os.path.join(TMP, 'hist_verify.json')

from PySide6.QtWidgets import QApplication, QFileDialog  # noqa: E402

app = QApplication(sys.argv)

import 发票助手Qt as M  # noqa: E402

CSV_OUT = os.path.join(TMP, 'test_export.csv')
QFileDialog.getSaveFileName = classmethod(
    lambda cls, *a, **k: (CSV_OUT, "CSV (*.csv)"))

M.SETTINGS.setValue("price_incl", "0")   # 固定默认口径: 不含税
M.SETTINGS.sync()
w = M.MainWindow()
w.show()
for _ in range(60):
    app.processEvents()
    time.sleep(0.01)

CLIP_DIR = TEST_DIR
import glob  # noqa: E402

w.add_paths(sorted(glob.glob(CLIP_DIR + r"\*.pdf")), join=False)
pump_early = [app.processEvents() or time.sleep(0.02) for _ in range(20)]

results = []


def check(name, cond, extra=""):
    results.append((name, bool(cond), extra))
    print(("PASS " if cond else "FAIL ") + name + (" | " + str(extra) if extra else ""))


def pump(n=30, dt=0.02):
    for _ in range(n):
        app.processEvents()
        time.sleep(dt)


check("导入文件=27", len(w.files) == 27, len(w.files))
check("文件表行数=27", w.file_table.rowCount() == 27, w.file_table.rowCount())

w.join_all()
pump()

# 等解析
t0 = time.time()
while time.time() - t0 < 90:
    if not any(f["parsing"] for f in w.files):
        break
    pump(5, 0.05)
check("全部加入=27", sum(1 for f in w.files if f["inReimb"]) == 27)
check("解析完成", not any(f["parsing"] for f in w.files))
check("解析失败=0", sum(1 for f in w.files if f["inv"] is None) == 0,
      sum(1 for f in w.files if f["inv"] is None))

tot = float(w.kpis["tTot"].text())
check("价税合计=1707.22", abs(tot - 1707.22) < 0.01, tot)
n_inv = int(w.kpis["tInv"].text())
check("发票张数=27", n_inv == 27, n_inv)

rows = w.rows_flat()
check("明细行=46", len(rows) == 46, len(rows))
check("明细表行数含合计", w.detail_table.rowCount() == len(rows) + 1,
      w.detail_table.rowCount())

amt = float(w.kpis["tAmt"].text())
tax = float(w.kpis["tTax"].text())
check("金额+税额≈1707.22", abs(amt + tax - 1707.22) < 0.05, (amt, tax))

# 移出/加入
first = w.files[0]["id"]
w.leave_by_id(first)
pump(15)
check("移出后=26", sum(1 for f in w.files if f["inReimb"]) == 26)
w.join_by_id(first)
pump(15)
check("再加入=27", sum(1 for f in w.files if f["inReimb"]) == 27)

# 按商品汇总
w.toggle_byitem()
pump(10)
n1 = w.detail_table.rowCount()
w.toggle_byitem()
pump(10)
n2 = w.detail_table.rowCount()
check("按商品汇总可切换", n1 > 0 and n2 == len(rows) + 1, (n1, n2))

# CSV
if os.path.exists(CSV_OUT):
    os.remove(CSV_OUT)
w.export_csv()
pump(10)
check("CSV 已生成", os.path.exists(CSV_OUT))
if os.path.exists(CSV_OUT):
    with io.open(CSV_OUT, encoding="utf-8-sig") as fh:
        lines = fh.read().strip().splitlines()
    check("CSV 行数≥47", len(lines) >= 47, len(lines))
    check("CSV 有商品摘要列", "商品摘要" in lines[0], lines[0][:60])

# 预览 + 缩放
w.preview(w.files[0])
pump(30, 0.03)
check("预览已加载", len(w.pdf_view._pages) >= 1, len(w.pdf_view._pages))
check("预览 hint 已清空", w.pdf_view.hint == "", repr(w.pdf_view.hint))

w.pdf_view.set_zoom(2.0)
pump(10)
check("缩放=2.0", abs(w.pdf_view._zoom - 2.0) < 1e-6, w.pdf_view._zoom)
w.pdf_view.fit_width()
pump(10)
check("适合宽度", abs(w.pdf_view._zoom - 1.0) < 1e-6, w.pdf_view._zoom)

# 空态提示
w.clear_preview()
pump(10)
check("清空后提示恢复", "单击左侧文件" in w.pdf_view.hint, repr(w.pdf_view.hint))

# 生成 Excel 报告
XLSX_OUT = CSV_OUT + ".xlsx"
if os.path.exists(XLSX_OUT):
    os.remove(XLSX_OUT)
w.generate_report()
pump(10)
check("报告 xlsx 已生成", os.path.exists(XLSX_OUT), XLSX_OUT)
if os.path.exists(XLSX_OUT):
    from openpyxl import load_workbook
    wb = load_workbook(XLSX_OUT)
    check("报告含4张表", len(wb.sheetnames) == 4, wb.sheetnames)
    check("报告汇总=1707.22",
          abs(float(wb["汇总"]["B7"].value) - 1707.22) < 0.01,
          wb["汇总"]["B7"].value)

# 价格含税口径(在清空前测)
w.set_price_incl(True)
pump(10)
tot_incl = float(w.kpis["tTot"].text())
amt_incl = float(w.kpis["tAmt"].text())
check("含税口径价税合计≈明细金额1603.66", abs(tot_incl - 1603.66) < 0.06,
      tot_incl)
check("含税后不含税金额<1603.66", amt_incl < 1603.66, amt_incl)
check("含税按钮显示✓", w.btn_incl.text().endswith("✓"), w.btn_incl.text())
w.set_price_incl(False)
pump(10)
check("切回不含税=1707.22",
      abs(float(w.kpis["tTot"].text()) - 1707.22) < 0.01,
      w.kpis["tTot"].text())
check("不含税按钮无✓", not w.btn_incl.text().endswith("✓"), w.btn_incl.text())

# 全部删除
M.QMessageBox.question = staticmethod(
    lambda *a, **k: M.QMessageBox.StandardButton.Yes)
w.remove_all()
pump(10)
check("全部删除后清空",
      len(w.files) == 0 and w.detail_table.rowCount() == 0
      and float(w.kpis["tTot"].text()) == 0.0)

print("=" * 44)
print("PASS %d / %d" % (sum(1 for _, c, _ in results if c), len(results)))
for n, c, e in results:
    if not c:
        print("FAILED:", n, e)
